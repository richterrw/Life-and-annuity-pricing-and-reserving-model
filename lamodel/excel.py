"""Formatted Excel workbook: executive summary first, then supporting sheets."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F3A5F")
SECTION_FONT = Font(bold=True, size=11, color="1F3A5F")
THIN = Side(style="thin", color="D9D9D9")

# number formats chosen by column name
FORMATS = [
    (("_bp",), "0.0"),
    (("a_x", "a-due"), "0.00000"),
    (("rate", "irr", "pct", "margin", "q_", "prob", "lapse_rate", "change_pct", "defer "), "0.00%"),
    (("per_1000", "per $1,000", "factor"), "#,##0.000"),
    (("inforce", "deaths", "lapses", "level"), "0.000000"),
    (("year", "age", "duration", "policies"), "0"),
]


def _fmt_for(col: str) -> str:
    c = str(col).lower()
    for keys, fmt in FORMATS:
        if any(k in c for k in keys):
            return fmt
    return "#,##0.00"


def _style_table(ws, df: pd.DataFrame, start_row: int = 1, pct_cols=None):
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=j)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        fmt = "0.00%" if pct_cols and col in pct_cols else _fmt_for(col)
        for i in range(start_row + 1, start_row + 1 + len(df)):
            c = ws.cell(row=i, column=j)
            if isinstance(c.value, (int, float)):
                c.number_format = fmt
            c.border = Border(bottom=THIN)
        width = max(len(str(col)), *(len(f"{v:,.4f}") if isinstance(v, float) else len(str(v)) for v in df[col])) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 10), 48)
    ws.row_dimensions[start_row].height = 30
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)


def write_workbook(
    path: Path, summary: list[tuple[str, list[tuple[str, object, str]]]], sheets: dict[str, pd.DataFrame]
):
    """``summary`` is a list of (section title, [(label, value, number_format)])."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame().to_excel(xw, sheet_name="Executive summary")
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
        wb = xw.book

        ws = wb["Executive summary"]
        ws["A1"] = "Life & Annuity Pricing and Reserving Model"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = "20-year level term and single-premium deferred income annuity. SOA 2017 CSO / 2012 IAM + Scale G2."
        ws["A2"].font = Font(italic=True, color="52514E")
        row = 4
        for title, items in summary:
            ws.cell(row=row, column=1, value=title).font = SECTION_FONT
            row += 1
            for label, value, fmt in items:
                ws.cell(row=row, column=1, value=label)
                c = ws.cell(row=row, column=2, value=value)
                if fmt:
                    c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
                ws.cell(row=row, column=1).border = ws.cell(row=row, column=2).border = Border(bottom=THIN)
                row += 1
            row += 1
        ws.column_dimensions["A"].width = 58
        ws.column_dimensions["B"].width = 22
        ws.sheet_view.showGridLines = False

        for name, df in sheets.items():
            _style_table(wb[name[:31]], df)
